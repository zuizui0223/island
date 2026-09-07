"""Drop only completely empty worksheets from a workbook copy.

The original source file is never modified.  This helper exists because some
published workbooks retain placeholder sheets that pandas cannot index after
read_excel returns an empty frame.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def sheet_is_empty(ws) -> bool:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                return False
    return True


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--input", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    args = p.parse_args()
    wb = load_workbook(args.input)
    dropped = []
    for ws in list(wb.worksheets):
        if sheet_is_empty(ws):
            dropped.append(ws.title)
            wb.remove(ws)
    if not wb.worksheets:
        raise ValueError("source workbook contains no non-empty worksheets")
    wb.save(args.output)
    print({"dropped_empty_sheets": dropped, "retained_sheets": wb.sheetnames})


if __name__ == "__main__":
    main()
